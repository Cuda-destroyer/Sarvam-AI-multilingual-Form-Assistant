"""
DOCX Form Reader & Field Extractor.
Parses Microsoft Word (.docx) documents, detecting form structures in tables,
paragraphs, underscore runs, and explicit placeholders.
"""

import io
from pathlib import Path
import re
from typing import BinaryIO, List, Tuple, Union
from docx import Document
from core.models import DestinationAnchor, FieldSpec, FieldType, HandlingMode
from core.privacy.logger import get_logger

logger = get_logger("docx_reader")

# Common patterns for form fields in paragraphs
UNDERSCORE_PATTERN = re.compile(r"([A-Za-z0-9\s/().]+?)[:：]?\s*([_]{3,}|\[\s*\]|\[[A-Za-z\s]+\])")
PLACEHOLDER_PATTERN = re.compile(r"\{\{([A-Za-z0-9_]+)\}\}")


def infer_field_type_from_label(label: str) -> Tuple[FieldType, HandlingMode, List[str]]:
    """Deterministically infer field type, handling mode, and options from label."""
    clean = label.strip().lower()
    options: List[str] = []

    if any(k in clean for k in ["aadhaar", "uid"]):
        return (FieldType.AADHAAR, HandlingMode.NORMALIZE, options)
    if any(k in clean for k in ["mobile", "phone", "contact no", "cell"]):
        return (FieldType.MOBILE, HandlingMode.NORMALIZE, options)
    if any(k in clean for k in ["dob", "date of birth", "birth date", "date"]):
        return (FieldType.DATE, HandlingMode.NORMALIZE, options)
    if any(k in clean for k in ["email", "e-mail"]):
        return (FieldType.EMAIL, HandlingMode.NORMALIZE, options)
    if any(k in clean for k in ["pin code", "pincode", "postal code", "zip"]):
        return (FieldType.NUMBER, HandlingMode.NORMALIZE, options)
    if any(k in clean for k in ["gender", "sex"]):
        return (FieldType.CHOICE, HandlingMode.COPY_EXACTLY, ["Male", "Female", "Other"])
    if any(k in clean for k in ["name", "applicant", "father", "mother", "guardian"]):
        return (FieldType.NAME, HandlingMode.TRANSLITERATE, options)
    if any(k in clean for k in ["address", "residence", "domicile"]):
        return (FieldType.ADDRESS, HandlingMode.TRANSLATE, options)

    return (FieldType.TEXT, HandlingMode.COPY_EXACTLY, options)


def slugify_label(label: str, index: int) -> str:
    """Generate a stable identifier from label."""
    clean = re.sub(r"[^\w\s]", "", label).strip().lower()
    clean = re.sub(r"\s+", "_", clean)
    return clean[:30] or f"field_{index}"


class DocxFormReader:
    """Extracts fillable form fields from DOCX files."""

    def extract_fields(self, doc_bytes_or_path: Union[bytes, Path, str, BinaryIO]) -> List[FieldSpec]:
        """
        Analyze DOCX document and return list of extracted FieldSpecs with exact anchors.
        """
        if isinstance(doc_bytes_or_path, (str, Path)):
            doc = Document(str(doc_bytes_or_path))
        elif isinstance(doc_bytes_or_path, bytes):
            doc = Document(io.BytesIO(doc_bytes_or_path))
        else:
            doc = Document(doc_bytes_or_path)

        fields: List[FieldSpec] = []
        field_counter = 0

        # 1. Inspect Tables (Most official institutional forms use tables)
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                cells = row.cells
                if len(cells) >= 2:
                    # Case A: Label cell followed by an empty or placeholder cell
                    label_text = cells[0].text.strip()
                    val_text = cells[1].text.strip()

                    # Check if cell[0] looks like a label and cell[1] is empty or has underscores
                    if label_text and (not val_text or "_" in val_text or "[" in val_text):
                        field_counter += 1
                        ftype, handling, opts = infer_field_type_from_label(label_text)
                        clean_label = label_text.rstrip(":").strip()
                        fid = slugify_label(clean_label, field_counter)

                        anchor = DestinationAnchor(
                            target_type="docx_cell",
                            table_index=t_idx,
                            row_index=r_idx,
                            col_index=1,
                        )
                        spec = FieldSpec(
                            id=fid,
                            label=clean_label,
                            field_type=ftype,
                            handling=handling,
                            options=opts,
                            anchor=anchor,
                            display_order=field_counter,
                        )
                        fields.append(spec)

                if len(cells) >= 4:
                    # Case B: 4-column layout: [Label1, Value1, Label2, Value2]
                    label2_text = cells[2].text.strip()
                    val2_text = cells[3].text.strip()
                    if label2_text and (not val2_text or "_" in val2_text or "[" in val2_text):
                        field_counter += 1
                        ftype, handling, opts = infer_field_type_from_label(label2_text)
                        clean_label = label2_text.rstrip(":").strip()
                        fid = slugify_label(clean_label, field_counter)

                        anchor = DestinationAnchor(
                            target_type="docx_cell",
                            table_index=t_idx,
                            row_index=r_idx,
                            col_index=3,
                        )
                        spec = FieldSpec(
                            id=fid,
                            label=clean_label,
                            field_type=ftype,
                            handling=handling,
                            options=opts,
                            anchor=anchor,
                            display_order=field_counter,
                        )
                        fields.append(spec)

        # 2. Inspect Paragraphs for underscore runs or placeholders
        for p_idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            # Check for {{field_id}} template tags
            placeholders = PLACEHOLDER_PATTERN.findall(text)
            for ph in placeholders:
                field_counter += 1
                clean_label = ph.replace("_", " ").title()
                ftype, handling, opts = infer_field_type_from_label(clean_label)
                anchor = DestinationAnchor(
                    target_type="docx_placeholder",
                    placeholder_text=f"{{{{{ph}}}}}",
                )
                fields.append(FieldSpec(
                    id=ph.lower(),
                    label=clean_label,
                    field_type=ftype,
                    handling=handling,
                    options=opts,
                    anchor=anchor,
                    display_order=field_counter,
                ))

            # Check for "Label: _________" patterns
            matches = UNDERSCORE_PATTERN.findall(text)
            for m in matches:
                label_candidate = m[0].strip()
                if 2 < len(label_candidate) < 50:
                    field_counter += 1
                    clean_label = label_candidate.rstrip(":").strip()
                    ftype, handling, opts = infer_field_type_from_label(clean_label)
                    fid = slugify_label(clean_label, field_counter)
                    anchor = DestinationAnchor(
                        target_type="docx_placeholder",
                        placeholder_text=m[1],
                    )
                    fields.append(FieldSpec(
                        id=fid,
                        label=clean_label,
                        field_type=ftype,
                        handling=handling,
                        options=opts,
                        anchor=anchor,
                        display_order=field_counter,
                    ))

        logger.info("DOCX reader extracted %d fields from document.", len(fields))
        return fields
