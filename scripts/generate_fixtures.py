"""
Generate test fixtures for Phase 0 and unit tests:
1. college_admission_form.docx (DOCX with tables and paragraphs)
2. scholarship_application_form.pdf (AcroForm PDF with fillable fields)
3. bank_account_opening.pdf (Flat text PDF with coordinate lines)
"""

from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def generate_docx_form(output_path: Path):
    doc = Document()
    doc.add_heading("National College of Engineering - Admission Form 2026", level=1)
    doc.add_paragraph("Please provide the candidate details in the form below.")

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"

    fields = [
        ("Applicant Name:", ""),
        ("Date of Birth:", ""),
        ("Mobile Number:", ""),
        ("Aadhaar Number:", ""),
        ("Email Address:", ""),
        ("Permanent Address:", ""),
    ]

    for idx, (label, val) in enumerate(fields):
        row = table.rows[idx]
        row.cells[0].text = label
        row.cells[1].text = val

    p = doc.add_paragraph()
    p.add_run("\nParent / Guardian Name: _______________________")
    p.add_run("\nPIN Code: _______________________")

    doc.save(str(output_path))
    print(f"Generated DOCX: {output_path}")


def generate_flat_pdf_form(output_path: Path):
    c = canvas.Canvas(str(output_path), pagesize=letter)
    width, height = letter

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, height - 50, "State Bank of India - Account Opening Application")

    c.setFont("Helvetica", 11)
    c.drawString(50, height - 80, "Kindly complete the form details in block letters.")

    form_items = [
        ("1. Full Name:", height - 130),
        ("2. Date of Birth:", height - 170),
        ("3. Mobile Number:", height - 210),
        ("4. Aadhaar Number:", height - 250),
        ("5. Email Address:", height - 290),
        ("6. Postal PIN Code:", height - 330),
    ]

    for label, y in form_items:
        c.drawString(50, y, label)
        c.drawString(180, y, "________________________________________")

    c.showPage()
    c.save()
    print(f"Generated Flat PDF: {output_path}")


def generate_acroform_pdf(output_path: Path):
    c = canvas.Canvas(str(output_path), pagesize=letter)
    width, height = letter

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, height - 50, "National Merit Scholarship Application Form")
    c.setFont("Helvetica", 11)
    c.drawString(50, height - 80, "Interactive Government Application Form")

    c.drawString(50, height - 130, "Applicant Name:")
    c.acroForm.textfield(name="applicant_name", x=160, y=height - 135, width=300, height=20)

    c.drawString(50, height - 170, "Date of Birth:")
    c.acroForm.textfield(name="date_of_birth", x=160, y=height - 175, width=200, height=20)

    c.drawString(50, height - 210, "Mobile Number:")
    c.acroForm.textfield(name="mobile_number", x=160, y=height - 215, width=200, height=20)

    c.drawString(50, height - 250, "Aadhaar Number:")
    c.acroForm.textfield(name="aadhaar_number", x=160, y=height - 255, width=250, height=20)

    c.drawString(50, height - 290, "Email Address:")
    c.acroForm.textfield(name="email_address", x=160, y=height - 295, width=300, height=20)

    c.showPage()
    c.save()
    print(f"Generated AcroForm PDF: {output_path}")


if __name__ == "__main__":
    fixtures_dir = Path("tests/fixtures/forms")
    fixtures_dir.mkdir(parents=True, exist_ok=True)
    generate_docx_form(fixtures_dir / "college_admission_form.docx")
    generate_flat_pdf_form(fixtures_dir / "bank_account_opening.pdf")
    generate_acroform_pdf(fixtures_dir / "scholarship_application_form.pdf")
